"""O extrato sintético tem o mesmo formato (e as mesmas esquisitices) do da B3."""

import hashlib
import zipfile
from datetime import date
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook

from sample_data.b3_format import (
    EMPTY,
    HEADER,
    Direction,
    Movement,
    StatementRow,
    write_movement_statement,
)

BROKER = "CORRETORA ALFA S/A"


def row(day: date, movement: Movement = Movement.LIQUIDACAO, **kwargs) -> StatementRow:
    defaults = {
        "direction": Direction.CREDIT,
        "date": day,
        "movement": movement,
        "product": "BBAS3 - BANCO DO BRASIL S/A",
        "institution": BROKER,
        "quantity": Decimal("100"),
        "unit_price": Decimal("25.40"),
        "amount": Decimal("2540.00"),
    }
    return StatementRow(**(defaults | kwargs))


def read_rows(path: Path) -> list[tuple]:
    return list(load_workbook(path).active.iter_rows(values_only=True))


def test_header_matches_b3(tmp_path: Path) -> None:
    path = write_movement_statement([row(date(2021, 3, 1))], tmp_path / "x.xlsx")
    assert read_rows(path)[0] == HEADER


def test_newest_first_and_same_day_keeps_order(tmp_path: Path) -> None:
    rows = [
        row(date(2021, 3, 1), product="A"),
        row(date(2021, 5, 1), product="B"),
        row(date(2021, 3, 1), product="C"),
    ]
    data = read_rows(write_movement_statement(rows, tmp_path / "x.xlsx"))[1:]
    assert [r[3] for r in data] == ["B", "A", "C"]


def test_date_is_text_and_missing_values_are_dash(tmp_path: Path) -> None:
    split = row(
        date(2021, 3, 1), Movement.DESDOBRO, quantity=Decimal("100"), unit_price=None, amount=None
    )
    (cells,) = read_rows(write_movement_statement([split], tmp_path / "x.xlsx"))[1:]
    assert cells[1] == "01/03/2021"
    assert cells[6] == EMPTY and cells[7] == EMPTY


def test_integers_stay_integers(tmp_path: Path) -> None:
    fraction = row(date(2021, 3, 1), quantity=Decimal("0.5"), unit_price=Decimal("10.00"))
    (cells,) = read_rows(write_movement_statement([fraction], tmp_path / "x.xlsx"))[1:]
    assert cells[5] == 0.5
    assert cells[6] == 10 and isinstance(cells[6], int)


def test_same_rows_same_bytes(tmp_path: Path) -> None:
    rows = [row(date(2021, 3, 1)), row(date(2022, 1, 10), Movement.DIVIDENDO)]
    first = write_movement_statement(rows, tmp_path / "a.xlsx").read_bytes()
    second = write_movement_statement(rows, tmp_path / "b.xlsx").read_bytes()
    assert hashlib.sha256(first).digest() == hashlib.sha256(second).digest()


def test_no_timestamp_of_generation_in_file(tmp_path: Path) -> None:
    # Dois arquivos gerados no mesmo segundo seriam iguais de qualquer jeito; o que
    # garante o determinismo é não haver data de geração em lugar nenhum.
    path = write_movement_statement([row(date(2021, 3, 1))], tmp_path / "x.xlsx")
    with zipfile.ZipFile(path) as z:
        assert {i.date_time for i in z.infolist()} == {(2020, 1, 1, 0, 0, 0)}
        core = z.read("docProps/core.xml").decode()
    assert core.count("2020-01-01T00:00:00Z") == 2
