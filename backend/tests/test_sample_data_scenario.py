"""A história sintética contém os cenários prometidos e o gabarito é coerente."""

from collections import Counter
from pathlib import Path

from openpyxl import load_workbook

from sample_data import expected
from sample_data.b3_format import Direction, Movement
from sample_data.generate import generate
from sample_data.scenario import BROKER_A, movement_rows


def test_has_identical_rows_on_same_day() -> None:
    counts = Counter(movement_rows())
    assert any(n >= 2 for n in counts.values())


def test_amount_is_quantity_times_price() -> None:
    for row in movement_rows():
        if row.unit_price is not None:
            assert row.amount == (row.quantity * row.unit_price).quantize(row.amount)


def test_uses_more_than_one_spelling_of_the_same_broker() -> None:
    used = {row.institution for row in movement_rows()}
    assert len(used & set(BROKER_A.spellings)) >= 2


def test_has_sales_with_profit_and_loss() -> None:
    results = [sale.result for sale in expected.SALES]
    assert any(r > 0 for r in results) and any(r < 0 for r in results)


def test_answer_key_adds_up() -> None:
    # Confere só a aritmética interna do gabarito, não o PM (isso é papel do motor).
    for sale in expected.SALES:
        assert sale.proceeds - sale.cost == sale.result
    for position in expected.POSITIONS.values():
        assert position.quantity * position.average_price == position.total_cost


def test_every_sale_in_answer_key_is_in_statement() -> None:
    # Venda de verdade (débito com preço) ou leilão de fração.
    sales = {
        (r.date, r.product.split(" - ")[0], r.quantity, r.amount)
        for r in movement_rows()
        if (r.direction is Direction.DEBIT and r.unit_price is not None)
        or r.movement is Movement.LEILAO_FRACAO
    }
    for sale in expected.SALES:
        assert (sale.date, sale.ticker, sale.quantity, sale.proceeds) in sales


def test_generate_writes_statement(tmp_path: Path) -> None:
    (path,) = generate(tmp_path)
    rows = list(load_workbook(path).active.iter_rows(values_only=True))
    assert len(rows) == 1 + len(movement_rows())


def test_every_loan_out_returns_later_with_same_quantity() -> None:
    # Saída de aluguel: mesmo rótulo da venda, mas sem preço.
    rows = movement_rows()
    outs = [
        r
        for r in rows
        if r.movement is Movement.LIQUIDACAO
        and r.direction is Direction.DEBIT
        and r.unit_price is None
    ]
    assert outs
    for out in outs:
        assert any(
            r.movement is Movement.LIQUIDACAO
            and r.direction is Direction.CREDIT
            and r.product == out.product
            and r.quantity == out.quantity
            and r.date > out.date
            for r in rows
        )


def test_loan_income_in_answer_key_is_in_statement() -> None:
    labels = {"REMUNERACAO": Movement.EMPRESTIMO, "REEMBOLSO": Movement.REEMBOLSO}
    found = {(r.date, r.product.split(" - ")[0], r.movement, r.amount) for r in movement_rows()}
    for income in expected.LOAN_INCOME:
        assert (income.date, income.ticker, labels[income.kind], income.amount) in found


def test_has_every_corporate_event_scenario() -> None:
    movements = {r.movement for r in movement_rows()}
    assert {
        Movement.DESDOBRO,
        Movement.GRUPAMENTO,
        Movement.FRACAO,
        Movement.LEILAO_FRACAO,
        Movement.BONIFICACAO,
        Movement.ATUALIZACAO,
        Movement.DIREITO,
        Movement.SUBSCRICAO_NAO_EXERCIDA,
    } <= movements


def test_events_carry_no_money() -> None:
    # Na B3, só o leilão de fração (entre os eventos) vem com preço e valor.
    no_money = {
        Movement.DESDOBRO,
        Movement.GRUPAMENTO,
        Movement.FRACAO,
        Movement.BONIFICACAO,
        Movement.ATUALIZACAO,
        Movement.DIREITO,
        Movement.SUBSCRICAO_NAO_EXERCIDA,
    }
    for r in movement_rows():
        if r.movement in no_money:
            assert r.unit_price is None and r.amount is None, r
