"""O consolidado sintético imita o da B3 e carrega a divergência proposital."""

from pathlib import Path

import pytest
from openpyxl import load_workbook

from sample_data import expected
from sample_data.generate import generate
from sample_data.scenario import consolidated_march_2024


@pytest.fixture(scope="module")
def workbook(tmp_path_factory: pytest.TempPathFactory):
    _, consolidated = generate(tmp_path_factory.mktemp("out"))
    return load_workbook(consolidated)


def test_sheets_in_b3_order(workbook) -> None:
    assert workbook.sheetnames == [
        "Posição - Ações",
        "Posição - Fundos",
        "Posição - Renda Fixa",
        "Proventos Recebidos",
        "Negociações",
    ]


def test_product_has_trailing_spaces(workbook) -> None:
    product = workbook["Posição - Ações"]["A2"].value
    assert product != product.rstrip()


def test_income_quantity_is_text(workbook) -> None:
    assert isinstance(workbook["Proventos Recebidos"]["E2"].value, str)


def test_sheet_ends_with_total(workbook) -> None:
    rows = list(workbook["Posição - Ações"].iter_rows(values_only=True))
    assert rows[-2][-1] == "Total"
    stocks = consolidated_march_2024().stocks
    assert rows[-1][-1] == pytest.approx(float(sum(s.value for s in stocks)))


def test_fixed_income_total_ignores_shifted_title(workbook) -> None:
    rows = list(workbook["Posição - Renda Fixa"].iter_rows(values_only=True))
    assert rows[-1][16] == float(expected.FIXED_INCOME_SHEET_TOTAL)
    titles = consolidated_march_2024().fixed_income
    assert sum(t.value for t in titles) == expected.FIXED_INCOME_TRUE_TOTAL


def test_reported_quantity_carries_the_divergence() -> None:
    reported = {s.ticker: s.quantity for s in consolidated_march_2024().stocks}
    for ticker, (calculated, informed) in expected.RECONCILIATION_DIVERGENCES.items():
        assert calculated != informed
        assert reported[ticker] == informed


def test_same_report_same_bytes(tmp_path: Path) -> None:
    _, first = generate(tmp_path / "a")
    _, second = generate(tmp_path / "b")
    assert first.read_bytes() == second.read_bytes()
