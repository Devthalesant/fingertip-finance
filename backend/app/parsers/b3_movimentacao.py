"""Leitor do extrato de movimentação da B3 (xlsx).

Só lê: cada linha vira um MovementRow com tipos de verdade (data, Decimal) e o hash que
identifica a linha entre importações (ADR 0001). O que a linha significa (compra,
aluguel, provento) é trabalho do classificador, não daqui.
"""

import hashlib
import warnings
from collections import Counter
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import load_workbook

from app.models.enums import DataSource, Direction

SOURCE = DataSource.B3_MOVIMENTACAO

HEADER = (
    "Entrada/Saída",
    "Data",
    "Movimentação",
    "Produto",
    "Instituição",
    "Quantidade",
    "Preço unitário",
    "Valor da Operação",
)
DIRECTIONS = {"Credito": Direction.IN, "Debito": Direction.OUT}
# A B3 grava valor ausente como o texto "-".
EMPTY = "-"
CENT = Decimal("0.01")


@dataclass(frozen=True)
class MovementRow:
    row_number: int  # linha da planilha (1 é o cabeçalho)
    direction: Direction
    date: date
    movement: str  # rótulo da B3, como veio
    product: str
    institution: str  # grafia original; a tabela de aliases resolve depois
    quantity: Decimal
    unit_price: Decimal | None
    amount: Decimal | None
    payload: dict[str, str | None]  # as células originais, para a linha bruta
    content_hash: str
    occurrence_index: int  # quantas linhas de mesmo conteúdo vieram antes no arquivo
    row_hash: str


def parse_movement_statement(path: Path) -> list[MovementRow]:
    """Lê a primeira aba do extrato, na ordem do arquivo (a mais nova primeiro)."""
    with warnings.catch_warnings():
        # O xlsx da B3 não tem estilo padrão e o openpyxl avisa a cada leitura.
        warnings.filterwarnings("ignore", message="Workbook contains no default style")
        workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        lines = workbook.worksheets[0].iter_rows(values_only=True)
        header = next(lines, ())
        if tuple(header[: len(HEADER)]) != HEADER:
            raise ValueError(f"{path.name} não parece um extrato de movimentação da B3")
        rows = []
        seen: Counter[str] = Counter()
        for number, cells in enumerate(lines, start=2):
            cells = tuple(cells[: len(HEADER)])
            if all(c is None or c == "" for c in cells):
                continue
            rows.append(_parse_row(number, cells, seen))
        return rows
    finally:
        workbook.close()


def movement_from_raw(
    row_number: int,
    payload: dict[str, str | None],
    content_hash: str,
    occurrence_index: int,
    row_hash: str,
) -> MovementRow:
    """Relê uma linha bruta salva no banco (refazer o ledger sem o arquivo original)."""
    cells = tuple(payload[name] for name in HEADER)
    return MovementRow(
        row_number=row_number,
        payload=dict(payload),
        content_hash=content_hash,
        occurrence_index=occurrence_index,
        row_hash=row_hash,
        **_fields(row_number, cells),
    )


def _parse_row(number: int, cells: tuple, seen: Counter[str]) -> MovementRow:
    parsed = _fields(number, cells)
    content_hash = _content_hash(parsed)
    occurrence = seen[content_hash]
    seen[content_hash] += 1
    row_hash = _sha256(f"{SOURCE}|{content_hash}|{occurrence}")
    return MovementRow(
        row_number=number,
        payload={name: _cell_text(value) for name, value in zip(HEADER, cells, strict=True)},
        content_hash=content_hash,
        occurrence_index=occurrence,
        row_hash=row_hash,
        **parsed,
    )


def _fields(number: int, cells: tuple) -> dict:
    direction, day, movement, product, institution, quantity, price, amount = cells
    try:
        parsed = {
            "direction": DIRECTIONS[direction],
            "date": datetime.strptime(day, "%d/%m/%Y").date(),
            "movement": movement.strip(),
            "product": product.strip(),
            "institution": institution.strip(),
            "quantity": _decimal(quantity),
            "unit_price": _decimal(price),
            "amount": _money(_decimal(amount)),
        }
    except (KeyError, TypeError, ValueError) as exc:
        raise ValueError(f"Linha {number} do extrato inválida: {exc}") from exc
    if parsed["quantity"] is None:
        raise ValueError(f"Linha {number} do extrato sem quantidade")
    return parsed


def _decimal(value: object) -> Decimal | None:
    if value is None or value == EMPTY:
        return None
    if isinstance(value, bool):
        raise ValueError(f"número esperado, veio {value!r}")
    if isinstance(value, int):
        return Decimal(value)
    if isinstance(value, float):
        # repr devolve o menor texto que volta ao mesmo float: 6.5 → "6.5", nunca 6.4999.
        return Decimal(repr(value))
    if isinstance(value, str):
        # O payload salvo no banco guarda os números como texto.
        try:
            return Decimal(value.strip())
        except InvalidOperation:
            pass
    raise ValueError(f"número esperado, veio {value!r}")


def _money(value: Decimal | None) -> Decimal | None:
    if value is None:
        return None
    cents = value.quantize(CENT)
    if cents != value:
        raise ValueError(f"valor com fração de centavos: {value}")
    return cents


def _cell_text(value: object) -> str | None:
    if value is None or isinstance(value, str):
        return value
    return repr(value) if isinstance(value, float) else str(value)


def _content_hash(parsed: dict) -> str:
    """Hash do conteúdo normalizado: o mesmo em qualquer export que traga a linha.

    Espaços repetidos não contam (a B3 varia a grafia) e 3000 vale o mesmo que 3000.00.
    """

    def norm(value: object) -> str:
        if value is None:
            return ""
        if isinstance(value, Decimal):
            return format(value.normalize(), "f")
        if isinstance(value, date):
            return value.isoformat()
        return " ".join(str(value).split())

    return _sha256("|".join(norm(v) for v in parsed.values()))


def _sha256(text: str) -> str:
    return hashlib.sha256(text.encode()).hexdigest()
